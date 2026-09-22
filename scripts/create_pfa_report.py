from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'artifacts' / 'SmartConfirm_Rapport_PFA_FR.docx'
LOGO = ROOT / 'dashboard' / 'public' / 'smartconfirm_logo_v2.png'
SHOT1 = Path(r'C:\Users\AMINE\AppData\Local\Temp\codex-clipboard-5e1571bd-8279-4fec-a5c4-c3b74a873226.png')
SHOT2 = Path(r'C:\Users\AMINE\AppData\Local\Temp\codex-clipboard-9330d52d-e0f4-49c6-8bb5-6f4f0a092f31.png')

doc = Document()
sec = doc.sections[0]
sec.top_margin, sec.bottom_margin = Inches(.65), Inches(.65)
sec.left_margin, sec.right_margin = Inches(.75), Inches(.75)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(6)
for name, size, color in [('Title', 32, '0B1324'), ('Heading 1', 21, '0B1324'), ('Heading 2', 15, '149447'), ('Heading 3', 12, '1F5F99')]:
    st = styles[name]; st.font.name='Aptos Display'; st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color); st.font.bold=True

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def footer(section):
    p=section.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('SmartConfirm — Rapport PFA   •   ')
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); p._p.append(fld)

footer(sec)

def page_break(): doc.add_page_break()
def h1(text): doc.add_heading(text, level=1)
def h2(text): doc.add_heading(text, level=2)
def p(text, bold_prefix=None):
    para=doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        para.add_run(bold_prefix).bold=True; para.add_run(text[len(bold_prefix):])
    else: para.add_run(text)
    return para
def bullets(items):
    for item in items: doc.add_paragraph(item, style='List Bullet')
def numbered(items):
    for item in items: doc.add_paragraph(item, style='List Number')
def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; set_cell_shading(c,'149447'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.color.rgb=RGBColor(255,255,255); r.bold=True
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): cells[i].text=str(v); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return t

# Couverture
if LOGO.exists(): doc.add_picture(str(LOGO), width=Inches(2.5))
doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
title=doc.add_paragraph(); title.alignment=WD_ALIGN_PARAGRAPH.CENTER; title.space_after=Pt(6)
r=title.add_run('SMARTCONFIRM'); r.bold=True; r.font.size=Pt(36); r.font.color.rgb=RGBColor.from_string('0B1324')
sub=doc.add_paragraph('Automatisation intelligente des confirmations de commandes via WhatsApp'); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER
sub.runs[0].font.size=Pt(20); sub.runs[0].font.bold=True; sub.runs[0].font.color.rgb=RGBColor.from_string('149447')
doc.add_paragraph('\nProjet de fin d’année\n', style=None).alignment=WD_ALIGN_PARAGRAPH.CENTER
meta=doc.add_table(rows=5, cols=2); meta.alignment=WD_TABLE_ALIGNMENT.CENTER; meta.style='Table Grid'
for i,(a,b) in enumerate([('Présenté par','Amine'),('Année universitaire','2025–2026'),('Établissement','À compléter'),('Filière','À compléter'),('Encadrant','À compléter')]): meta.cell(i,0).text=a; meta.cell(i,1).text=b; set_cell_shading(meta.cell(i,0),'E8F8EE')
doc.add_paragraph('\nPlateforme SaaS multi-boutiques • IA conversationnelle • WhatsApp • Support humain').alignment=WD_ALIGN_PARAGRAPH.CENTER

page_break(); h1('Résumé')
p('SmartConfirm est une plateforme SaaS destinée aux commerçants qui souhaitent automatiser leurs échanges WhatsApp liés aux commandes. Elle connecte les boutiques Shopify, WooCommerce et YouCan à des sessions WhatsApp gérées par Baileys. À la réception d’un événement e-commerce, la plateforme synchronise la commande, planifie les messages, dialogue avec le client et peut exécuter des actions contrôlées : rechercher un produit, préparer un panier, créer ou confirmer une commande, modifier certaines informations et notifier l’évolution de la livraison.')
p('L’originalité du projet réside dans l’association d’une intelligence artificielle conversationnelle à un système de tool calling sécurisé. Le modèle ne manipule jamais directement la base de données ni les API des boutiques. Il exprime une intention structurée ; le backend vérifie les droits, valide les arguments, demande une confirmation explicite pour les actions sensibles puis exécute l’outil adapté au fournisseur. Un agent humain peut prendre le contrôle de toute la conversation, consulter le contexte du client et rendre ensuite la main à l’IA.')
p('La solution comprend également l’authentification JWT, la gestion des rôles, des plans et quotas, les paiements Stripe et PayPal, les campagnes, l’audio, les modèles de messages, l’audit et les statistiques. Ce rapport présente les besoins, l’architecture, les choix techniques, la réalisation, les tests, les limites et les perspectives du projet.')
p('Mots-clés : WhatsApp, commerce électronique, NestJS, React, PostgreSQL, intelligence artificielle, tool calling, webhooks, SaaS.')
h2('Abstract')
p('SmartConfirm is a SaaS commerce automation platform connecting WhatsApp with Shopify, WooCommerce and YouCan. It combines normalized webhooks, an action-oriented AI agent, guarded tool execution and human handoff to make order conversations useful, safe and traceable.')

page_break(); h1('Table des matières')
for line in ['1. Introduction et problématique','2. Étude des besoins','3. Conception et architecture','4. Modèle de données','5. Réalisation fonctionnelle','6. Intelligence artificielle et outils','7. Intégrations e-commerce','8. Collaboration avec le support humain','9. Sécurité, abonnements et paiements','10. Tests et validation','11. Déploiement','12. Résultats, limites et perspectives','Conclusion générale']:
    doc.add_paragraph(line)

page_break(); h1('1. Introduction et problématique')
h2('1.1 Contexte')
p('Dans de nombreux marchés, WhatsApp est le canal privilégié entre les boutiques en ligne et leurs clients. Pourtant, le catalogue, les commandes et le suivi restent dans la plateforme e-commerce. Le support doit alors copier les données, rechercher la bonne commande et répéter les mêmes questions.')
h2('1.2 Problématique')
p('Comment concevoir une plateforme multi-boutiques capable de comprendre des messages naturels, d’exécuter des actions commerciales fiables et de transférer la conversation à un humain, tout en évitant les doublons, les erreurs d’autorisation et la perte de contexte ?')
h2('1.3 Solution proposée')
p('SmartConfirm centralise les connexions aux boutiques et à WhatsApp. Les événements entrants sont normalisés, stockés puis exploités par des workflows déterministes ou par un agent IA. Toute action à impact métier passe par un service contrôlé et traçable.')
h2('1.4 Objectifs')
bullets(['Réduire le temps consacré aux confirmations manuelles.','Synchroniser les commandes, produits, clients et statuts.','Permettre des conversations naturelles en français, arabe et darija.','Créer un mécanisme de handoff complet entre IA et agent humain.','Proposer une administration SaaS avec plans, quotas et paiements.','Garantir l’isolation des données et l’audit des actions.'])

page_break(); h1('2. Étude des besoins')
h2('2.1 Acteurs')
table(['Acteur','Responsabilités'],[
('Client final','Discute via WhatsApp, consulte les produits, confirme ou modifie une demande.'),
('Utilisateur commerçant','Connecte ses sessions et boutiques, configure les automations, suit les commandes.'),
('Agent humain','Prend le contrôle du chat, consulte le contexte et exécute les actions d’assistance.'),
('Administrateur','Gère utilisateurs, plans, paiements, paramètres globaux et journaux.'),
('Plateforme e-commerce','Émet les webhooks et reçoit les mises à jour de commandes.')])
h2('2.2 Besoins fonctionnels')
bullets(['Authentification, profil et autorisations par rôle.','Connexion WhatsApp Baileys avec état, QR et compte unique.','Installation et reconnexion des boutiques.','Import et synchronisation des produits et commandes.','Messages automatiques configurables par événement.','Agent IA avec recherche, création et mise à jour contrôlées.','Support humain enrichi par le contexte client.','Campagnes, contacts, modèles de messages, audio et médias.','Abonnements, quotas, renouvellement, annulation et historique de paiement.'])
h2('2.3 Besoins non fonctionnels')
bullets(['Sécurité des secrets et vérification des signatures de webhook.','Idempotence des événements et prévention des doubles messages.','Architecture extensible à de nouveaux fournisseurs.','Observabilité, logs et audit.','Interface responsive, multilingue et thème clair/sombre.'])

page_break(); h1('3. Conception et architecture')
h2('3.1 Vue générale')
p('La solution suit une architecture modulaire. Le frontend React consomme une API NestJS. Le backend orchestre l’authentification, les sessions WhatsApp, les intégrations, les messages et la facturation. PostgreSQL constitue la source de vérité et Redis peut compléter les files d’attente ou le cache lorsque le déploiement l’active.')
table(['Couche','Technologies','Rôle'],[
('Présentation','React, Vite, TypeScript','Dashboard utilisateur et administration.'),
('API','NestJS, TypeScript','Contrôleurs, validation, services, sécurité et documentation Swagger.'),
('Domaine commerce','Services SmartConfirm','Commandes, produits, conversations, outils et automations.'),
('Intégration','Provider Registry','Adaptateurs Shopify, WooCommerce et YouCan.'),
('Données','PostgreSQL, TypeORM','Persistance relationnelle et migrations.'),
('Messagerie','Baileys','Connexion et échange avec WhatsApp.'),
('IA et audio','Providers configurables','Conversation, tool calling, STT et TTS selon le plan.')])
h2('3.2 Registre de fournisseurs')
p('IntegrationProviderRegistry sélectionne l’adaptateur correspondant à la boutique. Les services métier utilisent une interface commune au lieu de coder des comportements spécifiques à Shopify dans les modules génériques. Cette séparation facilite l’ajout d’un nouveau fournisseur et réduit les duplications.')

page_break(); h1('4. Modèle de données')
p('Le modèle relationnel est organisé autour de l’utilisateur, de la session WhatsApp et de la boutique. Une boutique possède des produits et des commandes. Les conversations IA, les paniers temporaires, les exécutions d’outils, les reçus de messages et les avis sont rattachés aux entités commerciales concernées.')
table(['Groupe','Tables principales','Finalité'],[
('Identité','user_accounts, user_login_sessions','Comptes, rôles et sessions JWT.'),
('Commerce','stores, products, orders, product_reviews','Catalogue, commandes et avis vérifiés.'),
('Conversation','order_ai_conversations, customer_support_conversations','Mémoire et contrôle IA/humain.'),
('Orchestration','commerce_tool_executions, commerce_message_receipts','Audit des outils et idempotence.'),
('WhatsApp','sessions, messages, message_batches','Connexion, messages et envois groupés.'),
('SaaS','billing_plans, subscriptions, payment_transactions','Plans, accès et historique financier.'),
('Intégration','integration_connections, ingress_events','Connexions externes et événements normalisés.')])
p('Les identifiants externes sont conservés séparément des UUID internes. Cette distinction évite d’envoyer à WooCommerce un UUID SmartConfirm à la place d’un identifiant numérique de produit ou de variation.')

page_break(); h1('5. Réalisation fonctionnelle')
h2('5.1 Gestion des sessions WhatsApp')
p('L’utilisateur crée une session dans la limite de son plan. Un nom interne unique évite les collisions entre comptes. Si la session est déjà authentifiée, l’interface ne redemande pas inutilement le QR. Le numéro connecté est vérifié afin d’éviter les doublons.')
h2('5.2 Gestion des boutiques')
p('L’assistant d’ajout ne demande que les informations nécessaires au fournisseur et la session WhatsApp à associer. Les informations de profil et de politique sont ensuite récupérées depuis la boutique lorsqu’elles existent. La page de gestion centralise la reconnexion, les synchronisations, les événements et les modèles de messages.')
h2('5.3 Automations de messages')
p('Chaque événement peut être activé ou désactivé et lié à un modèle : nouvelle commande, confirmation, annulation, mise à jour, fulfillment, suivi transporteur et livraison. Un délai minimal et des fenêtres horaires limitent les collisions de messages.')
h2('5.4 Campagnes et contacts')
p('Les campagnes sélectionnent une session et une audience, permettent d’exclure des numéros, de composer un message ou choisir un modèle, puis suivent les statuts envoyé, échoué, en attente et ignoré.')

if SHOT1.exists():
    doc.add_picture(str(SHOT1), width=Inches(6.7)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    cap=doc.add_paragraph('Figure 1 — Interface de discussion et contexte commercial.'); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER

page_break(); h1('6. Intelligence artificielle et outils')
h2('6.1 Principe')
p('L’IA est utilisée pour comprendre la langue naturelle et planifier la prochaine action. Elle reçoit un contexte réduit : identité du magasin, politiques utiles, produits pertinents, commandes du numéro et derniers tours de conversation. Cette stratégie réduit les coûts et la consommation de tokens.')
h2('6.2 Tool calling')
p('Le modèle retourne soit une réponse, soit un appel d’outil structuré. CommerceToolService vérifie l’outil, les arguments, le propriétaire de la boutique et la cohérence des identifiants. Les mutations sensibles passent par une confirmation explicite du client.')
table(['Catégorie','Exemples d’outils'],[
('Catalogue','Rechercher un produit, lister les variantes, consulter stock et images.'),
('Commandes','Lister les commandes du téléphone, consulter le détail, créer une commande.'),
('Mise à jour','Modifier adresse ou statut, confirmer ou annuler selon les droits.'),
('Support','Préparer un résumé, envoyer un rappel, transférer à un humain.'),
('Après-vente','Générer un PDF de commandes, demander un avis, suivre la livraison.')])
h2('6.3 Garde-fous')
bullets(['Validation DTO et contrôle d’accès avant exécution.','Séparation des identifiants internes et externes.','Confirmation obligatoire pour une création ou modification.','Journal de chaque appel avec résultat et erreur.','Limites de tours, tokens, STT et TTS selon le plan.','Fallback humain en cas d’erreur du fournisseur IA.'])

page_break(); h1('7. Intégrations e-commerce')
table(['Fournisseur','Connexion','Événements','Actions'],[
('Shopify','OAuth application','Orders, products, fulfillment','Import, tags, mises à jour et webhooks signés.'),
('WooCommerce','URL et clés REST','Order/product webhooks','Création de commande, statuts, notes et avis.'),
('YouCan','OAuth','Rest hooks','Commandes, produits, clients et informations boutique.')])
h2('7.1 Pipeline de webhook')
numbered(['Vérifier la signature et l’horodatage.','Construire un IngressEvent normalisé.','Dédupliquer grâce à l’identifiant externe.','Mettre à jour les entités locales.','Exécuter les automations activées.','Enregistrer la livraison et le résultat.'])
h2('7.2 Création depuis WhatsApp')
p('Une commande créée par le bot est d’abord construite dans un panier local. Après confirmation et validation de l’adresse, le provider crée la commande distante. Le résultat est immédiatement persisté localement, sans attendre le webhook de retour. Lorsque ce webhook arrive, l’idempotence met à jour le même enregistrement sans envoyer un second message.')

page_break(); h1('8. Collaboration avec le support humain')
p('La page Chats devient un poste de travail. Pour chaque numéro, l’agent consulte la boutique, les commandes, les produits, les quantités, le paiement, le fulfillment, la confirmation, la dernière conversation IA et la prochaine action recommandée.')
bullets(['Take over : l’agent prend la main sur toute la conversation.','Resume AI : l’agent réactive l’automatisation avec le contexte sauvegardé.','Préparer un résumé de commande dans le composeur.','Envoyer un rappel de confirmation ou un média produit.','Créer une commande assistée depuis le catalogue de la boutique.','Identifier rapidement un blocage de paiement, d’adresse ou de stock.'])
if SHOT2.exists():
    doc.add_picture(str(SHOT2), width=Inches(6.7)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    cap=doc.add_paragraph('Figure 2 — Panneau de contexte destiné à l’agent humain.'); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER

page_break(); h1('9. Sécurité, abonnements et paiements')
h2('9.1 Sécurité')
bullets(['JWT pour les comptes utilisateurs et administration par rôle.','Isolation systématique par userId sur sessions, boutiques, messages et statistiques.','Chiffrement des secrets de boutiques et des clés fournisseurs.','Vérification HMAC ou signature selon le fournisseur.','Validation stricte des entrées et journal d’audit.','Le compte administrateur gère la plateforme sans utiliser les fonctions commerciales utilisateur.'])
h2('9.2 Plans et quotas')
p('Les plans sont stockés en base et peuvent définir le nombre de sessions, boutiques, messages envoyés et reçus, tokens IA, transcriptions, réponses audio et fonctionnalités avancées. Lorsqu’un abonnement expire ou qu’un quota est atteint, les actions concernées sont désactivées et l’interface propose une mise à niveau.')
h2('9.3 Paiements')
p('Stripe et PayPal créent les abonnements. Les webhooks de paiement sont la source de vérité : ils associent le client à l’utilisateur, enregistrent la transaction, actualisent la date de renouvellement et mettent à jour le statut de l’abonnement. Les changements de plan doivent être idempotents et gérer le prorata, l’annulation différée et les remboursements.')

page_break(); h1('10. Tests et validation')
table(['Niveau','Tests recommandés','Critères'],[
('Unitaire','Normalisation, quotas, téléphone, calcul du panier','Résultats déterministes et cas limites.'),
('Intégration','Providers, base, JWT, webhooks','Données isolées et erreurs traduites.'),
('E2E','OAuth, commande WhatsApp, paiement','Scénario complet sans double message.'),
('Sécurité','Signature, rôles, secrets, injections','Aucun accès inter-utilisateur.'),
('Charge','Campagnes, webhooks, files','Débit stable et reprise après erreur.'),
('UX','Mobile, thèmes, langues, formulaires','Parcours compréhensible et accessible.')])
h2('Scénario de démonstration')
numbered(['Se connecter avec un compte utilisateur.','Vérifier une session WhatsApp connectée.','Connecter ou ouvrir une boutique de test.','Demander un produit depuis WhatsApp.','Créer et confirmer une commande.','Vérifier la commande dans la boutique et dans SmartConfirm.','Prendre le contrôle dans Chats puis reprendre l’IA.'])

page_break(); h1('11. Déploiement')
p('Le projet est conteneurisé avec Docker. Les variables d’environnement configurent PostgreSQL, JWT, URLs publiques, fournisseurs IA, paiements et intégrations. En production, l’URL publique doit être utilisée pour Swagger, OAuth et webhooks au lieu de localhost.')
bullets(['PostgreSQL persistant et migrations au démarrage contrôlé.','Volume persistant pour les sessions Baileys et les plugins nécessaires.','HTTPS obligatoire pour OAuth et webhooks.','Health checks, logs structurés et métriques.','Sauvegarde de la base et rotation des secrets.','Workers séparés pour les campagnes et traitements longs lorsque la charge augmente.'])

page_break(); h1('12. Résultats, limites et perspectives')
h2('12.1 Résultats')
bullets(['Architecture SaaS multi-utilisateur et multi-boutiques.','Connexion WhatsApp et trois plateformes e-commerce.','Agent IA orienté actions et handoff humain.','Automations, campagnes, audio, notifications et avis.','Gestion des plans, quotas et paiements.'])
h2('12.2 Limites')
bullets(['La qualité dépend du modèle et de la langue audio.','Les API et règles des fournisseurs évoluent.','Les tunnels temporaires ne conviennent pas à la production.','La montée en charge exige des files et workers dédiés.','Les tests réels sur un grand volume restent nécessaires.'])
h2('12.3 Perspectives')
bullets(['Recherche augmentée sur politiques, FAQ et catalogue.','Analyse de sentiment et priorité automatique.','Plus de transporteurs et suivi unifié.','Observabilité métier et alertes proactives.','Application mobile pour les agents.','Évaluations automatiques de l’agent IA.'])

page_break(); h1('Conclusion générale')
p('SmartConfirm répond à un besoin concret : transformer WhatsApp en canal commercial relié à la source de vérité e-commerce. Le projet va au-delà d’un chatbot. Il combine intégrations, données, automatisation, intelligence artificielle, contrôle humain, sécurité et facturation dans une architecture extensible.')
p('Le principal enseignement est qu’une IA utile en production doit être encadrée par des outils métier, des validations et une traçabilité. Cette approche permet de conserver une expérience naturelle tout en garantissant que les actions sur les commandes restent cohérentes et réversibles.')
h2('Questions possibles pendant la soutenance')
table(['Question','Réponse courte'],[
('Pourquoi une base locale si la boutique possède déjà les commandes ?','Pour l’orchestration, la recherche rapide, l’audit et l’idempotence.'),
('Pourquoi le tool calling ?','Pour séparer le raisonnement linguistique de l’exécution sécurisée.'),
('Comment éviter les doubles messages ?','Identifiant d’événement, reçu de message et logique idempotente.'),
('Que se passe-t-il si l’IA échoue ?','Fallback contrôlé, journal d’erreur et transfert humain.'),
('Comment ajouter une plateforme ?','Implémenter l’interface provider puis l’enregistrer dans le registry.'),
('Pourquoi PostgreSQL ?','Relations fortes, transactions, requêtes analytiques et robustesse en production.')])

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
